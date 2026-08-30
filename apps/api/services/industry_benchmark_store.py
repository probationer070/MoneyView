"""Storage and parsing for Damodaran's industry-average vintages.

The vintage key is the dataset's PUBLICATION date, not the fetch date. The data
changes annually, so a fetch-dated row would manufacture variation that did not
occur -- the same argument
`docs/superpowers/specs/2026-07-28-statements-acquisition-and-manual-snapshots-design.md`
makes against daily snapshots of quarterly statements.

Loading is manual: `store_vintage(vintage, parse_workbook(path))`. Wiring this
into the acquisition layer's scheduler is deliberately not done -- an annual
dataset does not need one, and a scheduler for it would be machinery ahead of
need.
"""

from __future__ import annotations

import openpyxl

from apps.api.services.db import get_db
from apps.api.services.industry_maps import (
    EXCLUDED_ROWS,
    SECTOR_TO_INDUSTRIES,
    damodaran_industry_for_yahoo,
    sector_for_industry,
)
from packages.core_finance.industry_benchmark import (
    BENCHMARK_COLUMNS,
    BenchmarkUnavailable,
    IndustryRow,
    SectorBenchmark,
    resolve_benchmark,
)

_VALUE_COLUMNS = tuple(column.key for column in BENCHMARK_COLUMNS)


def parse_workbook(
    path: str, *, sheet: str = "Industry Average Beta (US)"
) -> list[IndustryRow]:
    """Read one vintage out of Damodaran's workbook.

    Columns are located by HEADER TEXT, not position: he republishes annually
    and column order is not a contract.
    """
    worksheet = openpyxl.load_workbook(path, data_only=True)[sheet]
    header_row = [cell.value for cell in worksheet[1]]
    index = {str(name).strip(): position for position, name in enumerate(header_row) if name}

    required = ["Industry Name", "Number of firms"] + [
        c.source_header for c in BENCHMARK_COLUMNS if c.required
    ]
    missing = [name for name in required if name not in index]
    if missing:
        raise ValueError(
            f"{path} sheet {sheet!r} is missing required headers: {missing}. "
            f"Found: {sorted(index)}"
        )

    rows: list[IndustryRow] = []
    for raw in worksheet.iter_rows(min_row=2, values_only=True):
        name = raw[index["Industry Name"]]
        firms = raw[index["Number of firms"]]
        if not isinstance(name, str) or not isinstance(firms, (int, float)):
            continue
        if name.strip() in EXCLUDED_ROWS:
            continue
        rows.append(IndustryRow(
            name=name.strip(),
            firms=int(firms),
            values={
                column.key: (
                    float(raw[index[column.source_header]])
                    if column.source_header in index
                    and isinstance(raw[index[column.source_header]], (int, float))
                    else None
                )
                for column in BENCHMARK_COLUMNS
            },
        ))
    return rows


def store_vintage(vintage: str, rows: list[IndustryRow]) -> int:
    """Persist one vintage, replacing any existing rows for the same key."""
    columns = ", ".join(_VALUE_COLUMNS)
    placeholders = ", ".join("?" * (len(_VALUE_COLUMNS) + 3))
    with get_db() as conn:
        conn.executemany(
            f"INSERT OR REPLACE INTO industry_benchmark "
            f"(vintage, industry_name, firms, {columns}) VALUES ({placeholders})",
            [
                (vintage, row.name, row.firms,
                 *(row.values.get(key) for key in _VALUE_COLUMNS))
                for row in rows
            ],
        )
    return len(rows)


def load_vintage(vintage: str) -> list[IndustryRow]:
    columns = ", ".join(_VALUE_COLUMNS)
    with get_db() as conn:
        found = conn.execute(
            f"SELECT industry_name, firms, {columns} FROM industry_benchmark "
            f"WHERE vintage = ? ORDER BY industry_name",
            (vintage,),
        ).fetchall()
    return [
        IndustryRow(
            name=row["industry_name"],
            firms=int(row["firms"]),
            values={key: row[key] for key in _VALUE_COLUMNS},
        )
        for row in found
    ]


def latest_vintage(on_or_before: str | None = None) -> str | None:
    """Newest stored vintage at or before `on_or_before`, or None."""
    with get_db() as conn:
        if on_or_before is None:
            row = conn.execute(
                "SELECT MAX(vintage) AS vintage FROM industry_benchmark"
            ).fetchone()
        else:
            row = conn.execute(
                "SELECT MAX(vintage) AS vintage FROM industry_benchmark "
                "WHERE vintage <= ?",
                (on_or_before,),
            ).fetchone()
    return row["vintage"] if row and row["vintage"] else None


def _stored_industry(ticker: str) -> str:
    with get_db() as conn:
        row = conn.execute(
            "SELECT industry FROM corporate_quote_facts WHERE ticker = ?",
            (ticker.upper(),),
        ).fetchone()
    return (row["industry"] or "") if row else ""


def resolve_for_ticker(
    ticker: str, *, as_of: str | None = None
) -> tuple[SectorBenchmark | None, str | None, str | None]:
    """Resolve a sector benchmark for one ticker.

    Returns `(benchmark, vintage, None)` or `(None, None, reason)`. Exactly one
    of benchmark/reason is non-None: a missing or unreliable benchmark yields NO
    case rather than a silently degraded one. Falling back to an all-industry
    average would produce a number that looks like a sector benchmark and is not
    one. The vintage travels with a successful resolution so callers -- notably
    `build_conservative_case`, which stamps the vintage into `case_name` and
    every `evidence_source` -- never have to call `latest_vintage()` a second
    time and risk it returning a different vintage than the one actually used.
    """
    vintage = latest_vintage(on_or_before=as_of)
    if vintage is None:
        return None, None, "no_vintage: no industry benchmark data has been loaded"

    industry = _stored_industry(ticker)
    if not industry:
        return None, None, f"no_industry: {ticker} has no industry from the quote source"

    mapped = damodaran_industry_for_yahoo(industry)
    if mapped is None:
        return None, None, (
            f"unmapped_industry: {ticker}'s industry {industry!r} is not in "
            f"YAHOO_TO_DAMODARAN; add it to apps/api/services/industry_maps.py"
        )

    sector = sector_for_industry(mapped)
    if sector is None:
        return None, None, (
            f"unmapped_industry: {mapped!r} is in no sector in SECTOR_TO_INDUSTRIES"
        )

    names = set(SECTOR_TO_INDUSTRIES[sector])
    rows = [row for row in load_vintage(vintage) if row.name in names]
    try:
        return resolve_benchmark(sector, rows), vintage, None
    except BenchmarkUnavailable as exc:
        return None, None, f"sector_too_thin: {exc}"
